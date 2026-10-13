from pathlib import Path
from collections import Counter
import csv
from datetime import datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from io import BytesIO
import re
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import load_workbook


# __file__ représente le fichier normalize_excel.py.
# parents[1] permet de remonter jusqu'au dossier Projet_final.
PROJECT_DIR = Path(__file__).resolve().parents[1]

SOURCE_DIR = PROJECT_DIR / "data" / "source_excel"
OUTPUT_DIR = PROJECT_DIR / "data" / "raw_csv"
REJECT_DIR = PROJECT_DIR / "data" / "rejects"

# Exemple reconnu : 000001000127500 (T)  02/01/2026  10:30
TICKET_PATTERN = re.compile(
    r"^(?P<ticket_id>\d+) \((?P<ticket_type>[A-Z])\)\s+"
    r"(?P<sale_date>\d{2}/\d{2}/\d{4})\s+"
    r"(?P<sale_time>\d{2}:\d{2})$"
)

PAYMENT_METHODS = {
    "CARTES",
    "ESPECES",
    "CHEQUES",
    "TICKET RESTAURANT",
}

MONEY_PRECISION = Decimal("0.01")

TICKET_COLUMNS = [
    "ticket_id",
    "ticket_type",
    "sale_timestamp",
    "customer_name",
    "ticket_total_ttc",
    "seller_name",
    "seller_id",
    "calculated_total_ttc",
    "total_difference",
    "article_count",
    "payment_methods",
    "source_month",
    "source_file",
    "source_row_number",
]

ARTICLE_COLUMNS = [
    "ticket_id",
    "ticket_type",
    "sale_timestamp",
    "customer_name",
    "seller_id",
    "product_code",
    "product_label_raw",
    "quantity",
    "unit_price_ttc",
    "discount_percent",
    "vat_rate",
    "gross_amount_ttc",
    "discount_amount_ttc",
    "net_amount_ttc",
    "source_month",
    "source_file",
    "source_row_number",
]

REJECT_COLUMNS = [
    "source_file",
    "source_row_number",
    "reason",
    "details",
]


def find_source_files():
    """Retourne les fichiers Excel mensuels triés par nom."""

    files = sorted(SOURCE_DIR.glob("????_??.xlsx"))

    if not files:
        raise FileNotFoundError(
            f"Aucun fichier mensuel trouvé dans {SOURCE_DIR}"
        )

    return files


def is_number(value):
    """Indique si une cellule contient un nombre, même stocké comme texte."""

    if value is None or isinstance(value, bool):
        return False

    try:
        Decimal(str(value).strip().replace(",", "."))
    except InvalidOperation:
        return False

    return True


def to_decimal(value):
    """Convertit une valeur Excel en nombre décimal exact."""

    return Decimal(str(value).strip().replace(",", "."))


def to_text(value):
    """Convertit une cellule en texte propre sans transformer None en 'None'."""

    if value is None:
        return ""

    return str(value).strip()


def classify_row(row):
    """Classe une ligne Excel sans encore transformer ses valeurs."""

    code = row[0]
    label = row[1]

    if isinstance(code, str) and TICKET_PATTERN.match(code.strip()):
        return "ticket"

    if code == "CODE":
        return "repeated_header"

    if label == "SOUS TOTAL":
        return "subtotal"

    if (
        is_number(code)
        and isinstance(label, str)
        and label.strip()
        and is_number(row[2])
        and is_number(row[3])
    ):
        return "article"

    if isinstance(label, str) and label.strip() in PAYMENT_METHODS:
        return "payment"

    if all(value is None for value in row):
        return "empty"

    return "other"


def load_source_workbook(source_file):
    """Ouvre un export Excel sans modifier le fichier d'origine.

    Le logiciel de caisse écrit par erreur l'attribut ``biltinId`` dans la
    feuille de styles. LibreOffice le tolère, mais openpyxl attend
    ``builtinId``. La correction est donc appliquée uniquement à une copie
    conservée en mémoire.
    """

    corrected_workbook = BytesIO()

    with ZipFile(source_file, "r") as source_archive:
        with ZipFile(corrected_workbook, "w", ZIP_DEFLATED) as target_archive:
            for member in source_archive.infolist():
                content = source_archive.read(member.filename)

                if member.filename == "xl/styles.xml":
                    content = content.replace(b"biltinId", b"builtinId")

                target_archive.writestr(member, content)

    corrected_workbook.seek(0)

    return load_workbook(
        corrected_workbook,
        read_only=True,
        data_only=True,
    )


def inspect_workbook(source_file):
    """Compte les différents types de lignes d'un classeur."""

    workbook = load_source_workbook(source_file)

    worksheet = workbook.active
    counts = Counter()

    for row in worksheet.iter_rows(values_only=True):
        counts[classify_row(row)] += 1

    workbook.close()
    return counts


def extract_first_ticket(source_file):
    """Extrait le premier ticket et ses articles pour contrôler le modèle."""

    workbook = load_source_workbook(source_file)
    worksheet = workbook.active

    ticket = None
    articles = []

    for source_row_number, row in enumerate(
        worksheet.iter_rows(values_only=True),
        start=1,
    ):
        row_type = classify_row(row)

        if row_type == "ticket":
            # Le deuxième ticket indique que le premier est terminé.
            if ticket is not None:
                break

            match = TICKET_PATTERN.match(row[0].strip())
            sale_timestamp = datetime.strptime(
                f"{match.group('sale_date')} {match.group('sale_time')}",
                "%d/%m/%Y %H:%M",
            )

            ticket = {
                "ticket_id": match.group("ticket_id"),
                "ticket_type": match.group("ticket_type"),
                "sale_timestamp": sale_timestamp,
                "customer_name": to_text(row[2]),
                "ticket_total_ttc": to_decimal(row[5]),
                "seller_name": to_text(row[7]),
                "seller_id": to_text(row[9]),
                "source_file": source_file.name,
                "source_row_number": source_row_number,
            }

        elif row_type == "article" and ticket is not None:
            quantity = to_decimal(row[2])
            unit_price_ttc = to_decimal(row[3])
            discount_percent = to_decimal(row[4] or 0)

            gross_amount_ttc = quantity * unit_price_ttc
            discount_amount_ttc = (
                gross_amount_ttc * discount_percent / Decimal("100")
            )
            net_amount_ttc = gross_amount_ttc - discount_amount_ttc

            articles.append(
                {
                    "ticket_id": ticket["ticket_id"],
                    "product_code": to_text(row[0]),
                    "product_label_raw": to_text(row[1]),
                    "quantity": quantity,
                    "unit_price_ttc": unit_price_ttc,
                    "discount_percent": discount_percent,
                    "vat_rate": to_decimal(row[5]),
                    "gross_amount_ttc": gross_amount_ttc.quantize(
                        MONEY_PRECISION,
                        rounding=ROUND_HALF_UP,
                    ),
                    "discount_amount_ttc": discount_amount_ttc.quantize(
                        MONEY_PRECISION,
                        rounding=ROUND_HALF_UP,
                    ),
                    "net_amount_ttc": net_amount_ttc.quantize(
                        MONEY_PRECISION,
                        rounding=ROUND_HALF_UP,
                    ),
                    "source_row_number": source_row_number,
                }
            )

    workbook.close()
    return ticket, articles


def extract_workbook(source_file):
    """Extrait et contrôle tous les tickets et articles d'un classeur."""

    workbook = load_source_workbook(source_file)
    worksheet = workbook.active

    source_month = datetime.strptime(source_file.stem, "%Y_%m").strftime(
        "%Y-%m"
    )

    tickets = []
    articles = []
    rejects = []
    ticket_ids = set()
    current_ticket = None
    current_ticket_articles = []
    current_payment_methods = []

    def finalize_current_ticket():
        """Termine les calculs du ticket avant de passer au suivant."""

        if current_ticket is None:
            return

        calculated_total = sum(
            (article["net_amount_ttc"] for article in current_ticket_articles),
            start=Decimal("0.00"),
        ).quantize(MONEY_PRECISION, rounding=ROUND_HALF_UP)

        total_difference = (
            calculated_total - current_ticket["ticket_total_ttc"]
        ).quantize(MONEY_PRECISION, rounding=ROUND_HALF_UP)

        current_ticket["calculated_total_ttc"] = calculated_total
        current_ticket["total_difference"] = total_difference
        current_ticket["article_count"] = len(current_ticket_articles)
        current_ticket["payment_methods"] = "|".join(
            dict.fromkeys(current_payment_methods)
        )

        if abs(total_difference) > Decimal("0.02"):
            rejects.append(
                {
                    "source_file": source_file.name,
                    "source_row_number": current_ticket[
                        "source_row_number"
                    ],
                    "reason": "ticket_total_mismatch",
                    "details": (
                        f"total_caisse={current_ticket['ticket_total_ttc']};"
                        f"total_calcule={calculated_total};"
                        f"ecart={total_difference}"
                    ),
                }
            )

        tickets.append(current_ticket.copy())

    for source_row_number, row in enumerate(
        worksheet.iter_rows(values_only=True),
        start=1,
    ):
        row_type = classify_row(row)

        if row_type == "ticket":
            finalize_current_ticket()

            match = TICKET_PATTERN.match(row[0].strip())
            ticket_id = match.group("ticket_id")
            sale_timestamp = datetime.strptime(
                f"{match.group('sale_date')} {match.group('sale_time')}",
                "%d/%m/%Y %H:%M",
            )

            if ticket_id in ticket_ids:
                rejects.append(
                    {
                        "source_file": source_file.name,
                        "source_row_number": source_row_number,
                        "reason": "duplicate_ticket_id",
                        "details": ticket_id,
                    }
                )

            if sale_timestamp.strftime("%Y-%m") != source_month:
                rejects.append(
                    {
                        "source_file": source_file.name,
                        "source_row_number": source_row_number,
                        "reason": "ticket_month_mismatch",
                        "details": sale_timestamp.isoformat(),
                    }
                )

            ticket_ids.add(ticket_id)
            current_ticket = {
                "ticket_id": ticket_id,
                "ticket_type": match.group("ticket_type"),
                "sale_timestamp": sale_timestamp,
                "customer_name": to_text(row[2]),
                "ticket_total_ttc": to_decimal(row[5]),
                "seller_name": to_text(row[7]),
                "seller_id": to_text(row[9]),
                "source_month": source_month,
                "source_file": source_file.name,
                "source_row_number": source_row_number,
            }
            current_ticket_articles = []
            current_payment_methods = []

        elif row_type == "article":
            if current_ticket is None:
                rejects.append(
                    {
                        "source_file": source_file.name,
                        "source_row_number": source_row_number,
                        "reason": "article_without_ticket",
                        "details": to_text(row[1]),
                    }
                )
                continue

            quantity = to_decimal(row[2])
            unit_price_ttc = to_decimal(row[3])
            discount_percent = to_decimal(row[4] or 0)
            gross_amount_ttc = quantity * unit_price_ttc
            discount_amount_ttc = (
                gross_amount_ttc * discount_percent / Decimal("100")
            )
            net_amount_ttc = gross_amount_ttc - discount_amount_ttc

            article = {
                "ticket_id": current_ticket["ticket_id"],
                "ticket_type": current_ticket["ticket_type"],
                "sale_timestamp": current_ticket["sale_timestamp"],
                "customer_name": current_ticket["customer_name"],
                "seller_id": current_ticket["seller_id"],
                "product_code": to_text(row[0]),
                "product_label_raw": to_text(row[1]),
                "quantity": quantity,
                "unit_price_ttc": unit_price_ttc,
                "discount_percent": discount_percent,
                "vat_rate": to_decimal(row[5]),
                "gross_amount_ttc": gross_amount_ttc.quantize(
                    MONEY_PRECISION,
                    rounding=ROUND_HALF_UP,
                ),
                "discount_amount_ttc": discount_amount_ttc.quantize(
                    MONEY_PRECISION,
                    rounding=ROUND_HALF_UP,
                ),
                "net_amount_ttc": net_amount_ttc.quantize(
                    MONEY_PRECISION,
                    rounding=ROUND_HALF_UP,
                ),
                "source_month": source_month,
                "source_file": source_file.name,
                "source_row_number": source_row_number,
            }
            articles.append(article)
            current_ticket_articles.append(article)

        elif row_type == "payment" and current_ticket is not None:
            current_payment_methods.append(to_text(row[1]))

    finalize_current_ticket()
    workbook.close()

    return tickets, articles, rejects


def serialize_csv_value(value):
    """Transforme une valeur Python en texte stable pour le CSV."""

    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")

    if isinstance(value, Decimal):
        return format(value, "f")

    if value is None:
        return ""

    return value


def write_csv(rows, columns, target_file):
    """Ecrit un CSV complet, puis le rend disponible en une seule operation."""

    target_file.parent.mkdir(parents=True, exist_ok=True)
    temporary_file = target_file.with_suffix(target_file.suffix + ".tmp")

    with temporary_file.open(
        "w",
        encoding="utf-8-sig",
        newline="",
    ) as csv_file:
        writer = csv.DictWriter(csv_file, fieldnames=columns)
        writer.writeheader()

        for row in rows:
            writer.writerow(
                {
                    column: serialize_csv_value(row.get(column))
                    for column in columns
                }
            )

    temporary_file.replace(target_file)


def write_monthly_outputs(source_file, tickets, articles, rejects):
    """Cree les trois fichiers CSV de controle pour un mois."""

    month = source_file.stem
    ticket_file = OUTPUT_DIR / f"tickets_{month}.csv"
    article_file = OUTPUT_DIR / f"sales_lines_{month}.csv"
    reject_file = REJECT_DIR / f"rejects_{month}.csv"

    write_csv(tickets, TICKET_COLUMNS, ticket_file)
    write_csv(articles, ARTICLE_COLUMNS, article_file)
    write_csv(rejects, REJECT_COLUMNS, reject_file)

    return ticket_file, article_file, reject_file


if __name__ == "__main__":
    source_files = find_source_files()

    print(f"{len(source_files)} fichier(s) trouvé(s) :")

    for source_file in source_files:
        counts = inspect_workbook(source_file)

        print(f"\n- {source_file.name}")
        print(f"  tickets  : {counts['ticket']}")
        print(f"  articles : {counts['article']}")
        print(f"  paiements: {counts['payment']}")
        print(f"  autres   : {counts['other']}")

    first_ticket, first_articles = extract_first_ticket(source_files[0])

    print("\nPremier ticket extrait :")
    print(f"  identifiant : {first_ticket['ticket_id']}")
    print(f"  date        : {first_ticket['sale_timestamp']}")
    print(f"  total caisse: {first_ticket['ticket_total_ttc']} EUR")
    print(f"  articles    : {len(first_articles)}")

    for article in first_articles:
        print(
            "  - "
            f"{article['product_label_raw']} | "
            f"quantité={article['quantity']} | "
            f"prix={article['unit_price_ttc']} | "
            f"net={article['net_amount_ttc']} EUR"
        )

    print("\nContrôle complet :")

    for source_file in source_files:
        tickets, articles, rejects = extract_workbook(source_file)
        ticket_file, article_file, reject_file = write_monthly_outputs(
            source_file,
            tickets,
            articles,
            rejects,
        )
        print(
            f"- {source_file.name}: "
            f"{len(tickets)} tickets, "
            f"{len(articles)} articles, "
            f"{len(rejects)} rejet(s)"
        )
        print(f"  cree : {ticket_file.name}")
        print(f"  cree : {article_file.name}")
        print(f"  cree : {reject_file.name}")
